from __future__ import annotations

import argparse
import io
import json
import os
import re
import secrets
from collections.abc import Callable
from pathlib import Path
from typing import Any

import httpx
from openpyxl import load_workbook
from playwright.sync_api import Page, sync_playwright

from browser_forms import AuditLog, fill_classified_form, page_state
from gliner_client import DecisionError, GlinerClient
from intake import run_clinic, run_helpdesk
from decision_clients import MODELS, make_client

RPA_URL = "https://rpachallenge.com/"
RPA_DATA_URL = "https://rpachallenge.com/assets/downloadFiles/challenge.xlsx"
PARABANK_URL = "https://parabank.parasoft.com/parabank/index.htm"
PIKO_URL = "https://pikoautomation.pl/sandbox/healing-agent"
PIKO_CHALLENGE_URL = "https://pikoautomation.pl/sandbox/automation-challenge/"
ROOT = Path(__file__).resolve().parent

# Chromium draws no text on some Linux setups when it loads every installed font (a bad font in the set
# breaks rendering: pages come back blank, so no form labels can be read). Limit it to DejaVu when that is
# installed; an explicit FONTCONFIG_FILE still wins.
if Path("/usr/share/fonts/truetype/dejavu").is_dir():
    os.environ.setdefault("FONTCONFIG_FILE", str(ROOT / "fonts" / "fontconfig.xml"))


def load_fixture(name: str) -> dict[str, str]:
    return json.loads((ROOT / "fixtures" / name).read_text(encoding="utf-8"))


def rpa_records() -> list[dict[str, str]]:
    response = httpx.get(RPA_DATA_URL, follow_redirects=True, timeout=30)
    response.raise_for_status()
    sheet = load_workbook(io.BytesIO(response.content), read_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    headers = {
        "First Name": "first_name",
        "Last Name": "last_name",
        "Company Name": "company_name",
        "Role in Company": "role_in_company",
        "Address": "address",
        "Email": "email",
        "Phone Number": "phone_number",
    }
    populated_columns = [
        index for index, value in enumerate(rows[0]) if value is not None
    ]
    keys = [headers[str(rows[0][index]).strip()] for index in populated_columns]
    return [
        {key: str(row[index]).strip() for key, index in zip(keys, populated_columns)}
        for row in rows[1:]
        if any(row[index] is not None for index in populated_columns)
    ]


def piko_challenge_records(workbook_path: str) -> list[dict[str, str]]:
    workbook = io.BytesIO(Path(workbook_path).read_bytes())
    sheet = load_workbook(workbook, read_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    headers = {
        "First Name": "first_name",
        "Last Name": "last_name",
        "Email": "email",
        "Phone": "phone",
        "Company": "company",
        "Job Title": "job_title",
        "Experience Level": "experience_level",
        "City": "city",
        "Country": "country",
    }
    keys = [headers[str(value).strip()] for value in rows[0]]
    return [
        {key: str(value).strip() for key, value in zip(keys, row)}
        for row in rows[1:]
        if any(value is not None for value in row)
    ]


RPA_DESCRIPTIONS = {
    "first_name": "Value for a field labeled First Name",
    "last_name": "Value for a field labeled Last Name",
    "company_name": "Value for a field labeled Company Name",
    "role_in_company": "Value for a field labeled Role in Company",
    "address": "Value for a field labeled Address",
    "email": "Value for a field labeled Email",
    "phone_number": "Value for a field labeled Phone Number",
}

PIKO_CHALLENGE_DESCRIPTIONS = {
    "first_name": "Value for a field labeled First Name",
    "last_name": "Value for a field labeled Last Name",
    "email": "Value for a field labeled Email",
    "phone": "Value for a field labeled Phone",
    "company": "Value for a field labeled Company",
    "job_title": "Value for a field labeled Job Title",
    "experience_level": "Value for a field labeled Experience Level",
    "city": "Value for a field labeled City",
    "country": "Value for a field labeled Country",
}


EventCallback = Callable[[dict[str, Any], bytes | None], None]


def run_rpa(
    page: Page, gliner: GlinerClient, on_event: EventCallback | None = None
) -> Path:
    audit = AuditLog("rpa-policyholder-intake", page=page, on_event=on_event)
    records = rpa_records()
    page.goto(RPA_URL, wait_until="networkidle")
    page.get_by_role("button", name="Start").click()
    # The form renders after Start; reading it too early finds no fields.
    page.locator("form input:visible").first.wait_for()
    for number, record in enumerate(records, start=1):
        fill_classified_form(
            page.locator("form"),
            record,
            RPA_DESCRIPTIONS,
            gliner,
            audit,
        )
        audit.write("record_ready", record=number)
        page.locator('input[type="submit"]').click()
    page.get_by_text(re.compile("congratulations", re.IGNORECASE)).wait_for()
    result = page_state(page)
    audit.write("completed", result=result)
    if "100%" not in result:
        raise RuntimeError(f"RPA Challenge did not report 100%: {result}")
    return audit.path


def run_piko_challenge(
    page: Page, gliner: GlinerClient, on_event: EventCallback | None = None
) -> Path:
    audit = AuditLog(
        "piko-excel-challenge",
        page=page,
        on_event=on_event,
        screenshot_events={"record_ready", "completed"},
    )
    page.goto(PIKO_CHALLENGE_URL, wait_until="networkidle")
    with page.expect_download() as download_info:
        page.get_by_role("link", name="Download the Excel file").click()
    download = download_info.value
    workbook_path = download.path()
    records = piko_challenge_records(str(workbook_path))
    audit.write(
        "file_downloaded", filename=download.suggested_filename, records=len(records)
    )

    page.get_by_role("button", name="Start", exact=True).click()
    audit.write("timer_started", records=len(records))
    for number, record in enumerate(records, start=1):
        fill_classified_form(
            page.locator(".ac-form-card"),
            record,
            PIKO_CHALLENGE_DESCRIPTIONS,
            gliner,
            audit,
        )
        audit.write("record_ready", record=number, total=len(records))
        page.get_by_role("button", name=re.compile(r"^Submit")).click()

    page.get_by_text(re.compile(r"10 / 10 correct records", re.IGNORECASE)).wait_for()
    result = page_state(page)
    audit.write("completed", result=result)
    if "100%" not in result:
        raise RuntimeError(f"Piko Automation Challenge did not report 100%: {result}")
    return audit.path


def synthetic_customer() -> dict[str, str]:
    suffix = secrets.token_hex(4)
    return {
        "first_name": "Avery",
        "last_name": "Morgan",
        "street_address": "17 Market Street",
        "city": "Austin",
        "state": "TX",
        "zip_code": "78701",
        "phone_number": "5550107",
        "ssn": f"900{secrets.randbelow(900000) + 100000}",
        "username": f"claimdemo{suffix}",
        "password": f"Demo-{suffix}-Safe",
    }


REGISTRATION_DESCRIPTIONS = {
    "first_name": "Value for the First Name field",
    "last_name": "Value for the Last Name field",
    "street_address": "Value for the Address or Street field",
    "city": "Value for the City field",
    "state": "Value for the State field",
    "zip_code": "Value for the Zip Code field",
    "phone_number": "Value for the Phone Number field",
    "ssn": "Value for the SSN field",
    "username": "Value for the Username field",
    "password": "Value for both Password and Confirm fields",
}


PAYMENT_DESCRIPTIONS = {
    "payee_name": "Value for the Payee Name field",
    "street_address": "Value for the Address or Street field",
    "city": "Value for the City field",
    "state": "Value for the State field",
    "zip_code": "Value for the Zip Code field",
    "phone_number": "Value for the Phone Number field",
    "claim_account": "Value for both Account and Verify Account fields",
    "settlement_amount": "Value for the payment Amount field",
}


def run_parabank(
    page: Page, gliner: GlinerClient, on_event: EventCallback | None = None
) -> Path:
    audit = AuditLog("parabank-claim-payment", page=page, on_event=on_event)
    customer = synthetic_customer()
    claim = load_fixture("claim.json")

    decision = gliner.classify(
        text=(
            "Insurance claim CLM-2026-1042 is approved. Fraud review is clear. "
            "All required documents are complete. Settlement amount is $25.00."
        ),
        task="claim_action",
        candidates={
            "issue_payment": "Approved claim with complete documents and clear fraud review",
            "request_documents": "Required claim documents are missing",
            "manual_review": "An adjuster must resolve ambiguity before payment",
            "refer_fraud": "Fraud indicators require investigation",
            "stop": "No browser action is safe",
        },
        prompt="Choose the permitted next insurance claim action.",
    )
    audit.write(
        "business_decision",
        claim_id=claim["claim_id"],
        decision=decision.label,
        confidence=decision.confidence,
        candidates=[
            "issue_payment",
            "request_documents",
            "manual_review",
            "refer_fraud",
            "stop",
        ],
    )
    if decision.label != "issue_payment":
        raise DecisionError(f"Claim was not approved for payment: {decision}")

    page.goto(PARABANK_URL, wait_until="domcontentloaded")
    page.get_by_role("link", name="Register").click()
    page.locator("#customerForm").wait_for()
    fill_classified_form(
        page.locator("#customerForm"),
        customer,
        REGISTRATION_DESCRIPTIONS,
        gliner,
        audit,
    )
    page.get_by_role("button", name="Register").click()
    page.get_by_text(
        re.compile("Your account was created successfully", re.IGNORECASE)
    ).wait_for()

    page.get_by_role("link", name="Bill Pay").click()
    page.get_by_role("button", name="Send Payment").wait_for()
    form = page.locator("form").filter(
        has=page.get_by_role("button", name="Send Payment")
    )
    fill_classified_form(
        form,
        {key: claim[key] for key in PAYMENT_DESCRIPTIONS},
        PAYMENT_DESCRIPTIONS,
        gliner,
        audit,
        ignored_labels={"From account #", "From Account"},
    )
    page.get_by_role("button", name="Send Payment").click()
    page.get_by_text(re.compile("Bill Payment Complete", re.IGNORECASE)).wait_for()
    audit.write("completed", claim_id=claim["claim_id"], result=page_state(page))
    return audit.path


PIKO_DESCRIPTIONS = {
    "first_name": "Value for the First Name field",
    "last_name": "Value for the Last Name field",
    "date_of_birth": "Value for the Date of Birth field",
    "gender": "Value for the Gender field",
    "street": "Value for the Street field",
    "house_number": "Value for the House or Apartment Number field",
    "city": "Value for the City field",
    "postal_code": "Value for the Postal Code field",
    "country": "Value for the Country field",
    "email": "Email address, not the preferred contact channel",
    "phone": "Value for the Phone field",
    "contact_preference": "Preferred contact channel (Email or Phone), not the email address",
}


def choose_page_action(
    page: Page,
    gliner: GlinerClient,
    audit: AuditLog,
    expected: str,
    description: str,
    candidates: dict[str, str],
) -> None:
    decision = gliner.classify(
        text=(
            f"Goal: complete an insured onboarding form. {description}\n"
            f"Current page state: {page_state(page)}"
        ),
        task="next_action",
        candidates=candidates,
        prompt="Choose exactly one permitted next browser action.",
    )
    audit.write(
        "browser_action",
        expected=expected,
        decision=decision.label,
        confidence=decision.confidence,
        candidates=list(candidates),
    )
    if decision.label != expected:
        raise DecisionError(f"Expected {expected}, model selected {decision}")


def run_piko(
    page: Page, gliner: GlinerClient, on_event: EventCallback | None = None
) -> Path:
    audit = AuditLog("piko-insured-onboarding", page=page, on_event=on_event)
    insured = load_fixture("insured.json")
    page.goto(PIKO_URL, wait_until="networkidle")
    page.get_by_role("switch", name="Toggle unstable UI mode").click()

    for step in range(3):
        card = page.locator(".sb-form-card")
        fill_classified_form(
            card,
            insured,
            PIKO_DESCRIPTIONS,
            gliner,
            audit,
        )
        action = "submit" if step == 2 else "continue"
        choose_page_action(
            page,
            gliner,
            audit,
            action,
            f"Step {step + 1} of 3 has been filled with validated synthetic data.",
            {
                action: (
                    "The current fields are complete; click Next to continue"
                    if action == "continue"
                    else "The final form step is complete; submit the form"
                ),
                "go_back": "Return to the previous step to correct information",
                "stop": "Stop because the current browser state is unsafe or unexpected",
            },
        )
        if action == "submit":
            page.get_by_role("button", name="Submit Form").click()
        else:
            page.get_by_role("button", name=re.compile("Next")).click()

        if step == 0:
            page.get_by_role("dialog").wait_for()
            choose_page_action(
                page,
                gliner,
                audit,
                "confirm_verification",
                "A session verification dialog blocks progress and asks for confirmation.",
                {
                    "confirm_verification": "Confirm verification and continue the form",
                    "dismiss": "Dismiss the dialog without continuing",
                    "stop": "Stop because the dialog is unsafe or unexpected",
                },
            )
            page.get_by_role("button", name="Confirm & Continue").click()

    page.get_by_text("Form submitted successfully").wait_for()
    audit.write("completed", result=page_state(page))
    return audit.path


DEMOS = {
    "rpa": run_rpa,
    "parabank": run_parabank,
    "piko": run_piko,
    "piko-excel": run_piko_challenge,
    "clinic": run_clinic,
    "helpdesk": run_helpdesk,
}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("demo", choices=[*DEMOS, "all"])
    parser.add_argument("--headed", action="store_true")
    parser.add_argument("--model", choices=sorted(MODELS), default="omni")
    args = parser.parse_args()

    gliner = make_client(args.model)
    names = list(DEMOS) if args.demo == "all" else [args.demo]
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=not args.headed)
        try:
            for name in names:
                context = browser.new_context(ignore_https_errors=True)
                page = context.new_page()
                try:
                    artifact = DEMOS[name](page, gliner)
                    print(f"{name}: passed, audit log: {artifact}")
                finally:
                    context.close()
        finally:
            browser.close()


if __name__ == "__main__":
    main()
